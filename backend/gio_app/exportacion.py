"""Reporte de folios para descarga en Excel o CSV.

El reporte respeta los mismos filtros, búsqueda y orden del tablero, y la visibilidad del rol,
pero sin paginación: incluye todos los folios que coinciden.
"""
import csv
import io
from datetime import datetime

from django.utils import timezone
from openpyxl import Workbook
from openpyxl.cell import WriteOnlyCell
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

MAX_FILAS = 50_000

# Excel y LibreOffice interpretan como fórmula una celda que empieza con estos caracteres.
INICIO_DE_FORMULA = ('=', '+', '-', '@', '\t', '\r')

TEXTO, FECHA, DECIMAL, ENTERO = 'texto', 'fecha', 'decimal', 'entero'


def _nombre(usuario):
    return usuario.nombre_completo if usuario else ''


def _telefono(inc):
    if inc.telefono_tecnico:
        return inc.telefono_tecnico
    return inc.tecnico.telefono if inc.tecnico else ''


def _local(fecha):
    return timezone.localtime(fecha).replace(tzinfo=None) if fecha else None


def _mttr(inc):
    return round(inc.mttr_horas, 2) if inc.mttr_horas is not None else None


# (encabezado, valor, tipo, ancho en caracteres)
COLUMNAS = [
    ('Folio', lambda i: i.folio, TEXTO, 14),
    ('Incidente', lambda i: i.incidente, TEXTO, 14),
    ('Empresa', lambda i: i.empresa, TEXTO, 30),
    ('Referencia', lambda i: i.referencia, TEXTO, 16),
    ('Tipo de servicio', lambda i: i.tipo_servicio, TEXTO, 16),
    ('Estatus', lambda i: i.get_estatus_display(), TEXTO, 12),
    ('Área operativa', lambda i: i.area_operativa, TEXTO, 14),
    ('COPE', lambda i: i.cope, TEXTO, 16),
    ('Central', lambda i: i.central, TEXTO, 22),
    ('Técnico de planta externa', lambda i: _nombre(i.tecnico), TEXTO, 26),
    ('Teléfono técnico', _telefono, TEXTO, 14),
    ('Técnico (planta interna)', lambda i: _nombre(i.evaluador), TEXTO, 26),
    ('Estado del enlace', lambda i: i.get_estado_enlace_display(), TEXTO, 14),
    ('Código de fallo', lambda i: i.codigo_fallo, TEXTO, 12),
    ('Dilación (días)', lambda i: i.dilacion_dias, ENTERO, 10),
    ('Apertura', lambda i: _local(i.fecha_apertura), FECHA, 17),
    ('Cierre', lambda i: _local(i.fecha_cierre), FECHA, 17),
    ('MTTR (horas)', _mttr, DECIMAL, 11),
    ('Diagnóstico final', lambda i: i.diagnostico_final, TEXTO, 40),
    ('Clave de liquidación', lambda i: i.cve_liq, TEXTO, 14),
    ('Descripción de liquidación', lambda i: i.desc_liq, TEXTO, 40),
    ('Evidencias', lambda i: i.num_evidencias, ENTERO, 10),
    ('Exportado a SISA', lambda i: 'Sí' if i.exportado_sisa else 'No', TEXTO, 10),
    ('Dirección', lambda i: i.dir_pta_a, TEXTO, 40),
    ('IP', lambda i: i.ips, TEXTO, 16),
]

FORMATO_FECHA = 'dd/mm/yyyy hh:mm'
FORMATO_DECIMAL = '0.00'


def neutralizar_formula(texto):
    """Antepone un apóstrofo a los textos que una hoja de cálculo ejecutaría como fórmula."""
    if isinstance(texto, str) and texto.startswith(INICIO_DE_FORMULA):
        return f"'{texto}"
    return texto


def nombre_de_archivo(extension):
    return f'GIO_Incidencias_{timezone.localtime():%Y%m%d_%H%M}.{extension}'


def generar_csv(incidentes):
    salida = io.StringIO()
    # BOM para que Excel abra el archivo en UTF-8 y respete los acentos.
    salida.write('﻿')
    escritor = csv.writer(salida, lineterminator='\r\n')
    escritor.writerow([encabezado for encabezado, *_ in COLUMNAS])
    for inc in incidentes:
        fila = []
        for _, valor, tipo, _ in COLUMNAS:
            dato = valor(inc)
            if dato is None:
                dato = ''
            elif tipo == FECHA:
                dato = dato.strftime('%d/%m/%Y %H:%M')
            elif tipo == TEXTO:
                dato = neutralizar_formula(dato)
            fila.append(dato)
        escritor.writerow(fila)
    return salida.getvalue().encode('utf-8')


def generar_xlsx(incidentes):
    libro = Workbook(write_only=True)
    hoja = libro.create_sheet('Incidencias')

    for indice, (_, _, _, ancho) in enumerate(COLUMNAS, start=1):
        hoja.column_dimensions[get_column_letter(indice)].width = ancho
    hoja.freeze_panes = 'A2'

    negrita = Font(bold=True, color='FFFFFF')
    relleno = PatternFill('solid', fgColor='1A2129')
    encabezados = []
    for encabezado, *_ in COLUMNAS:
        celda = WriteOnlyCell(hoja, value=encabezado)
        celda.font = negrita
        celda.fill = relleno
        celda.alignment = Alignment(vertical='center')
        encabezados.append(celda)
    hoja.append(encabezados)

    filas = 0
    for inc in incidentes:
        fila = []
        for _, valor, tipo, _ in COLUMNAS:
            dato = valor(inc)
            celda = WriteOnlyCell(hoja, value=dato if dato != '' else None)
            if tipo == FECHA and isinstance(dato, datetime):
                celda.number_format = FORMATO_FECHA
            elif tipo == DECIMAL:
                celda.number_format = FORMATO_DECIMAL
            elif tipo == TEXTO and isinstance(dato, str) and dato:
                # Se guarda siempre como texto: openpyxl convertiría en fórmula lo que empieza con "=".
                celda.data_type = 's'
            fila.append(celda)
        hoja.append(fila)
        filas += 1

    hoja.auto_filter.ref = f'A1:{get_column_letter(len(COLUMNAS))}{filas + 1}'

    salida = io.BytesIO()
    libro.save(salida)
    return salida.getvalue()
